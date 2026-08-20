from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart

wb = load_workbook("quarterly_summary.xlsx")
ws = wb["Summary"]

chart = BarChart()
chart.title = "Revenue by Quarter"
chart.x_axis.title = "Quarter"
chart.y_axis.title = "Revenue ($)"

data = Reference(ws, min_col=2, min_row=1, max_row=5)
categories = Reference(ws, min_col=1, min_row=2, max_row=5)

chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)

ws.add_chart(chart, "D2")
wb.save("quarterly_summary.xlsx")

line_chart = LineChart()
line_chart.title = "Revenue Trend by Quarter"
line_chart.x_axis.title = "Quarter"
line_chart.y_axis.title = "Revenue ($)"

line_chart.add_data(data, titles_from_data=True)
line_chart.set_categories(categories)

ws.add_chart(line_chart, "D18")
wb.save("quarterly_summary.xlsx")