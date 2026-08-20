from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import PatternFill

wb = load_workbook("sales_report.xlsx")
ws = wb.active

# 1. Formulas
for row in range(2, 10):
    ws[f"D{row}"] = f"=B{row}*C{row}"
ws["D10"] = "=SUM(D2:D9)"

# 2. Conditional formatting
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
ws.conditional_formatting.add(
    "B2:B9", CellIsRule(operator="lessThan", formula=["10"], fill=red_fill)
)
ws.conditional_formatting.add(
    "D2:D9", ColorScaleRule(start_type="min", start_color="F8696B",
                            end_type="max", end_color="63BE7B")
)

# 3. Freeze panes
ws.freeze_panes = "A2"

wb.save("sales_report.xlsx")