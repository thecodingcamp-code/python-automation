import pandas as pd
from openpyxl import load_workbook

# Step 1: pandas does the analysis
df = pd.read_excel("inventory.xlsx", sheet_name="Stock")
totals = df.groupby("Warehouse")["Quantity"].sum()

# Step 2: openpyxl writes the result into an existing, formatted report
wb = load_workbook("warehouse_report.xlsx")
ws = wb["Summary"]

ws["B2"] = "East"
ws["C2"] = int(totals["East"])
ws["B3"] = "West"
ws["C3"] = int(totals["West"])

wb.save("warehouse_report.xlsx")